"""Reporting services: CSV/Excel export, email subscriptions, data alerts and in-app notifications."""
from __future__ import annotations

import csv
import html
import io
import json
import logging
import smtplib
import ssl
from datetime import datetime, timezone
from email.message import EmailMessage
from typing import TYPE_CHECKING

from croniter import croniter

from .errors import UserError, public_message
from .notifications import notify
from .orchestrator import next_run
from .store import now, parse_ts

if TYPE_CHECKING:
    from .platform import Platform

log = logging.getLogger("tessera.reporting")
ALERT_SOURCES = ("query", "model", "quality", "freshness", "pipeline")
ALERT_OPS = {">": lambda a, b: a > b, ">=": lambda a, b: a >= b, "<": lambda a, b: a < b, "<=": lambda a, b: a <= b,
             "=": lambda a, b: a == b, "!=": lambda a, b: a != b}


class ReportError(UserError):
    pass


def fmt_value(v, field: dict | None) -> str:
    if v is None:
        return ""
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return str(v)
    f = (field or {}).get("format", "number")
    d = (field or {}).get("decimals", 2)
    if f == "currency":
        return f"${v:,.{d}f}"
    if f == "percent":
        return f"{v * 100 if abs(v) <= 1.5 else v:,.{d}f}%"
    if f == "integer":
        return f"{v:,.0f}"
    return f"{v:,.{d}f}" if isinstance(v, float) else f"{v:,}"


# ---------------------------------------------------------------- exports
def to_csv(res: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(res["columns"])
    w.writerows(res["rows"])
    return buf.getvalue().encode("utf-8-sig")


def _sheet_name(name: str, used: set[str]) -> str:
    base = "".join(ch for ch in name if ch not in '[]:*?/\\')[:28] or "Sheet"
    n, i = base, 2
    while n.lower() in used:
        n = f"{base[:25]} {i}"
        i += 1
    used.add(n.lower())
    return n


def to_xlsx(sheets: list[tuple[str, dict]], title: str = "Tessera export") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    head_fill = PatternFill("solid", fgColor="E3E9ED")
    for name, res in sheets:
        ws = wb.create_sheet(_sheet_name(name, used))
        fields = {f["name"]: f for f in res.get("fields") or []}
        labels = [fields.get(c, {}).get("label", c) for c in res["columns"]]
        ws.append(labels)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True), head_fill, Alignment(vertical="center")
        for r in res["rows"]:
            ws.append(r)
        for i, col in enumerate(res["columns"], start=1):
            f = fields.get(col, {})
            numfmt = {"currency": '"$"#,##0.00', "percent": "0.0%", "integer": "#,##0", "number": "#,##0.00"}.get(f.get("format"))
            width = max([len(str(labels[i - 1]))] + [len(str(r[i - 1])) for r in res["rows"][:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(9, width))
            if numfmt:
                for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    if isinstance(cell[0].value, (int, float)):
                        cell[0].number_format = numfmt
        ws.freeze_panes = "A2"
    if not sheets:
        wb.create_sheet("Empty")
    wb.properties.title = title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dashboard_results(p: Platform, user: dict, dashboard: dict, filters: list | None = None) -> list[tuple[str, dict, dict]]:
    """Run every query tile of a dashboard as `user`. Returns (title, query, result-or-error)."""
    out = []
    for t in dashboard.get("tiles") or []:
        if t.get("type") == "text" or not t.get("query_id"):
            continue
        q = p.store.one("SELECT * FROM queries WHERE id=?", (t["query_id"],))
        if not q:
            continue
        try:
            from . import workspaces as ws
            ws.require(p.store, user, q.get("workspace_id") or 1, "viewer")
            res = p.run_saved(user, q, {"filters": filters or []})
        except Exception as e:  # noqa: BLE001
            res = {"error": public_message(e), "columns": [], "rows": []}
        out.append((t.get("title") or q["name"], q, res))
    return out


def dashboard_email_html(p: Platform, dashboard: dict, results, link: str | None) -> str:
    parts = ["<div style=\"font-family:Segoe UI,Arial,sans-serif;color:#14212B;max-width:760px\">",
             f"<h2 style=\"margin:0 0 4px\">{html.escape(dashboard['name'])}</h2>",
             f"<p style=\"color:#5A6873;margin:0 0 16px\">{html.escape(dashboard.get('description') or '')} · {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}</p>"]
    for title, q, res in results:
        parts.append(f"<h3 style=\"margin:18px 0 6px;font-size:15px\">{html.escape(title)}</h3>")
        if res.get("error"):
            parts.append(f"<p style=\"color:#B8362F\">{html.escape(res['error'])}</p>")
            continue
        fields = {f["name"]: f for f in res.get("fields") or []}
        viz = (q.get("viz") or {}).get("type", "table")
        if viz == "kpi" and res["rows"]:
            col = res["columns"][-1]
            parts.append(f"<div style=\"font-size:28px;font-family:Consolas,monospace\">{html.escape((q['viz'].get('prefix') or '') + fmt_value(res['rows'][0][-1], fields.get(col)) + (q['viz'].get('suffix') or ''))}</div>")
            continue
        rows = res["rows"][:15]
        parts.append("<table style=\"border-collapse:collapse;font-size:13px\"><tr>" + "".join(
            f"<th style=\"text-align:left;padding:4px 10px;background:#E3E9ED\">{html.escape(fields.get(c, {}).get('label', c))}</th>" for c in res["columns"]) + "</tr>")
        for r in rows:
            parts.append("<tr>" + "".join(f"<td style=\"padding:4px 10px;border-bottom:1px solid #D2DAE0\">{html.escape(fmt_value(v, fields.get(c)))}</td>" for c, v in zip(res["columns"], r)) + "</tr>")
        parts.append("</table>")
        if len(res["rows"]) > 15:
            parts.append(f"<p style=\"color:#5A6873;font-size:12px\">{len(res['rows']) - 15} more rows in the attachment.</p>")
    if link:
        parts.append(f"<p style=\"margin-top:20px\"><a href=\"{html.escape(link)}\">Open the live dashboard</a></p>")
    parts.append("<p style=\"color:#5A6873;font-size:12px;margin-top:24px\">Sent by Tessera. Data reflects the subscription owner's access.</p></div>")
    return "".join(parts)


# ---------------------------------------------------------------- mail
class Mailer:
    def __init__(self, p: Platform):
        self.p = p

    def config(self) -> dict | None:
        c = self.p.store.get_setting("smtp")
        return c if c and c.get("host") else None

    def send(self, to: list[str], subject: str, html_body: str, attachments: list[tuple[str, bytes, str]] = ()) -> None:
        c = self.config()
        if not c:
            raise ReportError("Email is not configured. An admin can add SMTP settings under Admin → Email.")
        secret = self.p.vault.decrypt(c.get("secret_blob")).get("password")
        msg = EmailMessage()
        msg["Subject"], msg["From"], msg["To"] = subject, c.get("from") or c.get("username") or "tessera@localhost", ", ".join(to)
        msg.set_content("This message needs an HTML-capable mail client.")
        msg.add_alternative(html_body, subtype="html")
        for name, data, mime in attachments:
            maintype, subtype = mime.split("/", 1)
            msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
        port = int(c.get("port") or (465 if c.get("ssl") else 587))
        try:
            if c.get("ssl"):
                srv = smtplib.SMTP_SSL(c["host"], port, timeout=30, context=ssl.create_default_context())
            else:
                srv = smtplib.SMTP(c["host"], port, timeout=30)
                if c.get("starttls", True):
                    srv.starttls(context=ssl.create_default_context())
            with srv:
                if c.get("username"):
                    srv.login(c["username"], secret or "")
                srv.send_message(msg)
        except (smtplib.SMTPException, OSError) as e:
            raise ReportError(f"Email failed: {e}") from None


# ---------------------------------------------------------------- reporter
class Reporter:
    _last_purge: datetime | None = None
    purge = None

    def __init__(self, p: Platform, http_factory=None):
        self.p, self.mailer = p, Mailer(p)
        self.http_factory = http_factory

    def _http(self):
        import httpx
        return self.http_factory() if self.http_factory else httpx.Client(timeout=15)

    def _next(self, cron: str, at: datetime) -> str:
        return next_run(cron, at)

    def link(self, path: str) -> str | None:
        base = self.p.settings.base_url
        return f"{base}/app#{path}" if base else None

    # ---------------- subscriptions
    def send_subscription(self, sub: dict) -> dict:
        p = self.p
        owner = p.store.one("SELECT * FROM users WHERE id=? AND active=1", (sub["owner_id"],))
        d = p.store.one("SELECT * FROM dashboards WHERE id=?", (sub["dashboard_id"],))
        if not owner or not d:
            raise ReportError("The subscription's owner or dashboard no longer exists.")
        results = dashboard_results(p, owner, d, sub.get("filters") or [])
        body = dashboard_email_html(p, d, results, self.link(f"/dashboards/{d['id']}"))
        if sub.get("include_brief"):
            from . import brief
            try:
                eb = brief.dashboard_brief(p, owner, d, sub.get("filters") or [])
                if eb["headlines"]:
                    i = body.find("<h3")
                    body = body[:i] + brief.email_html(eb) + body[i:] if i >= 0 else body + brief.email_html(eb)
            except Exception as e:  # noqa: BLE001
                log.warning("executive brief for subscription %s failed: %s", sub.get("id"), e)
        atts = []
        stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        good = [(t, r) for t, _, r in results if not r.get("error")]
        if sub.get("attach") == "xlsx":
            atts.append((f"{d['name']} {stamp}.xlsx", to_xlsx(good, d["name"]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        elif sub.get("attach") == "csv":
            for t, r in good:
                atts.append((f"{t} {stamp}.csv", to_csv(r), "text/csv"))
        self.mailer.send(sub["recipients"], f"{d['name']} · {stamp}", body, atts)
        return {"recipients": len(sub["recipients"]), "tiles": len(results), "attachments": len(atts)}

    def run_subscription(self, sub: dict) -> dict:
        try:
            out = self.send_subscription(sub)
            status, detail = "sent", out
        except Exception as e:  # noqa: BLE001
            status, detail = "failed", {"error": public_message(e)}
            notify(self.p, sub["owner_id"], f"Subscription '{sub['name']}' failed", str(e), "#/dashboards/" + str(sub["dashboard_id"]), "bad")
        self.p.store.update("subscriptions", sub["id"], {"last_sent_at": now(), "last_status": status if status == "sent" else detail["error"][:300]})
        self.p.store.insert("deliveries", {"kind": "subscription", "ref_id": sub["id"], "ts": now(), "status": status, "detail": json.dumps(detail)})
        return {"status": status, **detail}

    # ---------------- alerts
    def alert_value(self, a: dict, owner: dict):
        src = a["source"]
        t = src.get("type")
        if t == "query":
            q = self.p.store.one("SELECT * FROM queries WHERE id=?", (src.get("query_id"),))
            if not q:
                raise ReportError("The alert's query was deleted.")
            from . import workspaces as ws
            ws.require(self.p.store, owner, q.get("workspace_id") or 1, "viewer")
            res = self.p.run_saved(owner, q, {"filters": src.get("filters") or []})
            col = src.get("field") or next((c for i, c in enumerate(res["columns"]) if any(isinstance(r[i], (int, float)) for r in res["rows"])), None)
            if not res["rows"] or col not in res["columns"]:
                return None
            return res["rows"][0][res["columns"].index(col)]
        if t == "model":
            m = self.p.model(src["model_id"])
            from . import workspaces as ws
            ws.require(self.p.store, owner, m.get("workspace_id") or 1, "viewer")
            res = self.p.run_model(owner, m, {"measures": [src["measure"]], "filters": src.get("filters") or []})
            return res["rows"][0][0] if res["rows"] else None
        if t == "quality":
            q = self.p.latest_quality(int(src["dataset_id"]))
            return None if q["score"] is None else round(q["score"] * 100, 2)
        if t == "freshness":
            d = self.p.dataset_by_id(int(src["dataset_id"]))
            lt = parse_ts(d.get("last_loaded_at"))
            return None if not lt else round((datetime.now(timezone.utc) - lt).total_seconds() / 3600, 2)
        if t == "pipeline":
            r = self.p.store.one("SELECT status FROM pipeline_runs WHERE pipeline_id=? AND status NOT IN ('queued','running') ORDER BY id DESC LIMIT 1",
                                 (src["pipeline_id"],))
            return 1 if r and r["status"] == "failed" else 0
        raise ReportError(f"Unknown alert source '{t}'")

    def evaluate(self, a: dict) -> dict:
        p = self.p
        owner = p.store.one("SELECT * FROM users WHERE id=? AND active=1", (a["owner_id"],))
        cond = a["condition"]
        try:
            if not owner:
                raise ReportError("Alert owner is deactivated.")
            v = self.alert_value(a, owner)
            if v is None:
                firing = bool(cond.get("fire_on_empty"))
            else:
                op = cond.get("op", ">")
                if a["source"].get("type") == "pipeline":
                    firing = v == 1
                else:
                    firing = ALERT_OPS[op](float(v), float(cond.get("threshold", 0)))
            new_state, err = ("firing" if firing else "ok"), None
        except Exception as e:  # noqa: BLE001
            v, new_state, err = None, "error", public_message(e)
        prev = a["state"]
        p.store.update("alerts", a["id"], {"state": new_state, "last_value": json.dumps(v), "last_checked": now(), "last_error": err})
        changed = new_state != prev and not (prev == "unknown" and new_state == "ok")
        if changed and new_state in ("firing", "ok", "error"):
            self._deliver(a, owner, new_state, v, err)
        return {"state": new_state, "value": v, "error": err, "notified": changed}

    def _deliver(self, a, owner, state, value, err):
        cond = a["condition"]
        if state == "firing":
            title = f"Alert: {a['name']}"
            body = f"Value {value} {cond.get('op', '')} {cond.get('threshold', '')}".strip() if a["source"].get("type") != "pipeline" else "The latest run failed."
            level = "bad"
        elif state == "ok":
            title, body, level = f"Resolved: {a['name']}", f"Value is now {value}.", "ok"
        else:
            title, body, level = f"Alert could not be checked: {a['name']}", err or "", "warn"
        ch = a["channels"] or {}
        link = "#/alerts"
        if ch.get("inapp", True) and owner:
            notify(self.p, owner["id"], title, body, link, level)
        results = {}
        if ch.get("email"):
            try:
                self.mailer.send(ch["email"], title, f"<div style=\"font-family:Segoe UI,Arial,sans-serif\"><h2>{html.escape(title)}</h2><p>{html.escape(body)}</p>"
                                 f"<p style=\"color:#5A6873\">Checked {now()} by Tessera.</p></div>")
                results["email"] = "sent"
            except ReportError as e:
                results["email"] = str(e)
        if ch.get("webhook"):
            try:
                with self._http() as cl:
                    r = cl.post(ch["webhook"], json={"event": f"alert.{state}", "alert": a["name"], "alert_id": a["id"], "value": value,
                                                     "condition": cond, "error": err, "at": now()})
                results["webhook"] = r.status_code
            except Exception as e:  # noqa: BLE001
                results["webhook"] = str(e)
        self.p.store.insert("deliveries", {"kind": "alert", "ref_id": a["id"], "ts": now(), "status": state, "detail": json.dumps({"value": value, **results})})

    # ---------------- scheduling
    def tick(self, at: datetime | None = None) -> dict:
        at = at or datetime.now(timezone.utc)
        done = {"alerts": 0, "subscriptions": 0}
        for a in self.p.store.all("SELECT * FROM alerts WHERE enabled=1"):
            due = parse_ts(a["next_run_at"])
            if due is None or due <= at:
                self.p.store.update("alerts", a["id"], {"next_run_at": self._next(a["schedule"], at)})
                if due is not None:
                    self.evaluate(a)
                    done["alerts"] += 1
        for s in self.p.store.all("SELECT * FROM subscriptions WHERE enabled=1"):
            due = parse_ts(s["next_run_at"])
            if due is None or due <= at:
                self.p.store.update("subscriptions", s["id"], {"next_run_at": self._next(s["schedule"], at)})
                if due is not None:
                    self.run_subscription(s)
                    done["subscriptions"] += 1
        purge = getattr(self, "purge", None)
        if purge and (self._last_purge is None or (at - self._last_purge).total_seconds() >= 86400):
            self._last_purge = at
            try:
                done["purged"] = purge()
            except Exception:  # noqa: BLE001
                log.exception("retention purge failed")
        return done

    def on_run_finished(self, pipeline: dict, run_id: int, status: str, who: str | None) -> None:
        if status == "failed":
            targets = {u["id"] for u in self.p.store.all("SELECT id FROM users WHERE email IN (?, ?) AND active=1",
                                                          (pipeline.get("created_by") or "", who or ""))}
            for uid in targets:
                notify(self.p, uid, f"Pipeline '{pipeline['name']}' failed", f"Run #{run_id} failed.", f"#/pipelines/run/{run_id}", "bad")
        for a in self.p.store.all("SELECT * FROM alerts WHERE enabled=1"):
            if a["source"].get("type") == "pipeline" and int(a["source"].get("pipeline_id", 0)) == pipeline["id"]:
                self.evaluate(a)


def validate_cron(expr: str) -> None:
    if not croniter.is_valid(expr or ""):
        raise ReportError(f"'{expr}' is not a valid cron expression (minute hour day month weekday).")


