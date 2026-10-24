"""Exports, email subscriptions (real SMTP), alerts and notifications."""
import email
import io
import json
import socket
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from conftest import Client, add_user, do_setup, make_app
from openpyxl import load_workbook


class Inbox:
    def __init__(self):
        self.messages = []

    async def handle_DATA(self, server, session, envelope):
        self.messages.append({"to": envelope.rcpt_tos, "msg": email.message_from_bytes(envelope.content)})
        return "250 OK"


@pytest.fixture
def smtp():
    aiosmtpd = pytest.importorskip("aiosmtpd.controller")
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    inbox = Inbox()
    ctl = aiosmtpd.Controller(inbox, hostname="127.0.0.1", port=port)
    ctl.start()
    yield inbox, port
    ctl.stop()


@pytest.fixture
def env(settings):
    hooks = []

    def handler(req):
        hooks.append(json.loads(req.content))
        return httpx.Response(200)
    app = make_app(settings, http_factory=lambda: httpx.Client(transport=httpx.MockTransport(handler)))
    c = Client(app)
    do_setup(c, settings, sample=True)
    return {"c": c, "app": app, "hooks": hooks}


def test_exports_xlsx_and_csv(env):
    c = env["c"]
    qid = next(q["id"] for q in c.get("/api/queries").json() if q["name"] == "Category performance")
    r = c.post(f"/api/queries/{qid}/export", json={"format": "xlsx"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]][:3] == ["Product category", "Orders", "Units sold"]
    assert ws.max_row == 7 and ws["D2"].number_format.startswith('"$"')
    csv = c.post(f"/api/queries/{qid}/export", json={"format": "csv"}).content.decode("utf-8-sig")
    assert csv.splitlines()[0].startswith("category,orders")
    dash = load_workbook(io.BytesIO(c.post("/api/dashboards/1/export", json={}).content))
    assert len(dash.sheetnames) == 9 and "Revenue by month" in dash.sheetnames
    sqlx = c.post("/api/sql/export", json={"sql": "select 1 as a, 'x' as b", "format": "csv"}).content.decode("utf-8-sig")
    assert sqlx.splitlines() == ["a,b", "1,x"]


def test_dashboard_filters_apply_to_export(env):
    c = env["c"]
    mid = c.get("/api/models").json()[0]["id"]
    qid = next(q["id"] for q in c.get("/api/queries").json() if q["name"] == "Revenue by province")
    r = c.post(f"/api/queries/{qid}/export", json={"format": "csv", "filters": [{"model_id": mid, "field": "province", "op": "in", "value": ["BC"]}]})
    lines = r.content.decode("utf-8-sig").splitlines()
    assert len(lines) == 2 and lines[1].startswith("BC,")


def test_subscription_sends_real_email_with_attachment(env, smtp):
    c = env["c"]
    inbox, port = smtp
    assert c.post("/api/subscriptions", json={"name": "Weekly sales", "dashboard_id": 1, "recipients": "a@acme.ca, b@acme.ca", "schedule": "0 8 * * 1"}).status_code == 200
    sid = c.get("/api/subscriptions").json()[0]["id"]
    r = c.post(f"/api/subscriptions/{sid}/send")
    assert r.status_code == 400 and "not configured" in r.json()["detail"]
    c.put("/api/admin/email", json={"host": "127.0.0.1", "port": port, "from": "tessera@acme.ca", "starttls": False})
    r = c.post(f"/api/subscriptions/{sid}/send")
    assert r.status_code == 200, r.text
    msg = inbox.messages[-1]
    assert set(msg["to"]) == {"a@acme.ca", "b@acme.ca"}
    assert msg["msg"]["Subject"].startswith("Sales overview")
    parts = list(msg["msg"].walk())
    html = next(p for p in parts if p.get_content_type() == "text/html").get_payload(decode=True).decode()
    assert "Revenue" in html and "$" in html
    att = next(p for p in parts if p.get_filename() and p.get_filename().endswith(".xlsx"))
    assert len(load_workbook(io.BytesIO(att.get_payload(decode=True))).sheetnames) == 9
    assert c.get("/api/subscriptions").json()[0]["last_status"] == "sent"
    assert c.post("/api/admin/email/test", json={"to": "ops@acme.ca"}).status_code == 200


def test_subscription_scheduling(env, smtp):
    c, app = env["c"], env["app"]
    inbox, port = smtp
    c.put("/api/admin/email", json={"host": "127.0.0.1", "port": port, "starttls": False})
    c.post("/api/subscriptions", json={"name": "Hourly", "dashboard_id": 1, "recipients": ["x@acme.ca"], "schedule": "0 * * * *", "attach": "csv"})
    rep = app.state.reporter
    assert rep.tick()["subscriptions"] == 0
    assert rep.tick(datetime.now(timezone.utc) + timedelta(hours=2))["subscriptions"] == 1
    att = [p.get_filename() for p in inbox.messages[-1]["msg"].walk() if p.get_filename()]
    assert len(att) == 9 and all(a.endswith(".csv") for a in att)


def test_subscription_validation(env):
    c = env["c"]
    assert c.post("/api/subscriptions", json={"dashboard_id": 1, "recipients": "not-an-email"}).status_code == 400
    assert c.post("/api/subscriptions", json={"dashboard_id": 1, "recipients": "a@b.ca", "schedule": "whenever"}).status_code == 400
    assert c.post("/api/subscriptions", json={"dashboard_id": 999, "recipients": "a@b.ca"}).status_code == 400


def test_alert_fires_resolves_and_notifies(env):
    c, hooks = env["c"], env["hooks"]
    mid = c.get("/api/models").json()[0]["id"]
    r = c.post("/api/alerts", json={"name": "Revenue too high", "source": {"type": "model", "model_id": mid, "measure": "revenue"},
                                    "condition": {"op": ">", "threshold": 1}, "channels": {"inapp": True, "webhook": "http://hooks.test/a"}})
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["state"] == "firing" and float(json.loads(a["last_value"])) > 1
    assert hooks[-1]["event"] == "alert.firing"
    n = c.get("/api/notifications").json()
    assert n["unread"] >= 1 and n["items"][0]["title"] == "Alert: Revenue too high"
    c.put(f"/api/alerts/{a['id']}", json={"condition": {"op": ">", "threshold": 1e12}})
    res = c.post(f"/api/alerts/{a['id']}/check").json()
    assert res["state"] == "ok" and res["notified"]
    assert hooks[-1]["event"] == "alert.ok"
    assert c.get("/api/notifications").json()["items"][0]["title"].startswith("Resolved")
    c.post("/api/notifications/read", json={})
    assert c.get("/api/notifications").json()["unread"] == 0
    again = c.post(f"/api/alerts/{a['id']}/check").json()
    assert not again["notified"], "no repeat notification while the state is unchanged"


def test_quality_freshness_and_pipeline_alerts(env):
    c, app = env["c"], env["app"]
    did = next(d["id"] for d in c.get("/api/datasets").json() if d["schema_name"] == "curated")
    r = c.post("/api/alerts", json={"name": "Quality under 99.9", "source": {"type": "quality", "dataset_id": did}, "condition": {"op": "<", "threshold": 99.9}}).json()
    assert r["state"] == "firing"
    r = c.post("/api/alerts", json={"name": "Stale", "source": {"type": "freshness", "dataset_id": did}, "condition": {"op": ">", "threshold": 24}}).json()
    assert r["state"] == "ok"
    p = c.post("/api/pipelines", json={"name": "doomed", "steps": [{"id": "x", "type": "transform", "sql": "select * from raw.nope", "target": "curated.z"}]}).json()
    al = c.post("/api/alerts", json={"name": "Doomed failed", "source": {"type": "pipeline", "pipeline_id": p["id"]}}).json()
    assert al["state"] == "ok"
    app.state.orch.trigger(p["id"], "manual", "admin@acme.ca", wait=True)
    after = next(x for x in c.get("/api/alerts").json() if x["id"] == al["id"])
    assert after["state"] == "firing", "pipeline alerts are evaluated as soon as the run finishes"
    titles = [n["title"] for n in c.get("/api/notifications").json()["items"]]
    assert "Pipeline 'doomed' failed" in titles


def test_alert_permissions(env):
    c, app = env["c"], env["app"]
    a = add_user(c, app, "al@acme.ca", "analyst")
    did = next(d["id"] for d in c.get("/api/datasets").json() if d["schema_name"] == "curated")
    assert a.post("/api/alerts", json={"name": "q", "source": {"type": "quality", "dataset_id": did}, "condition": {"op": "<", "threshold": 1}}).status_code == 403
    mine = c.post("/api/alerts", json={"name": "admin's", "source": {"type": "pipeline", "pipeline_id": 1}}).json()
    r = a.delete(f"/api/alerts/{mine['id']}")
    assert r.status_code == 404
    assert a.get("/api/alerts").json() == []
